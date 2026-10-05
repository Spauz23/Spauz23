import math, os, re, json, urllib.parse, urllib.request
ORIGIN="Egied De Jonghestraat 165, 2880 Bornem, Belgium"
CAP={"Bornem":"2880","Niel":"2845","Schelle":"2627","Puurs-Sint-Amands":"2870","Hemiksem":"2620","Boom":"2850","Temse":"9140","Willebroek":"2830","Kruibeke":"9150","Aartselaar":"2630","Rumst":"2840"}
def dest(ind,com):
    base=com.split(" (")[0]
    ind=re.sub(r"\s*\(.*?\)","",ind)
    return f"{ind}, {CAP.get(base,'')} {base}, Belgium"
def maps_link(d):
    return "https://www.google.com/maps/dir/?api=1&"+urllib.parse.urlencode({"origin":ORIGIN,"destination":d,"travelmode":"driving"})
KEY=os.environ.get("GOOGLE_MAPS_API_KEY")
def maps_dm(dests):
    # Distance Matrix: max 25 destinazioni per richiesta
    out=[]
    for i in range(0,len(dests),25):
        q=urllib.parse.urlencode({"origins":ORIGIN,"destinations":"|".join(dests[i:i+25]),"mode":"driving","units":"metric","key":KEY})
        r=json.load(urllib.request.urlopen("https://maps.googleapis.com/maps/api/distancematrix/json?"+q,timeout=30))
        for el in r["rows"][0]["elements"]:
            out.append((round(el["distance"]["value"]/1000,1),round(el["duration"]["value"]/60)) if el.get("status")=="OK" else (None,None))
    return out
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

O_LAT, O_LON = 51.1095, 4.2935  # Egied De Jonghestraat 165, Hingene (stima)

# (azienda, settore, indirizzo, comune, lat, lon, dipendenti, nota_dip, email, email_hr, tel)
D = [
("BMW Group Belux (BMW Belgium Luxembourg)","Automotive – sede Benelux","Lodderstraat 16","Bornem (Eikevliet)",51.0935,4.2830,170,"170 FTE (bilancio); ~600 dal 2027 con nuova sede","contact.be@bmw.be","","+32 3 890 50 02"),
("Topfloor","Pavimenti/arredo – produzione","Lodderstraat 18","Bornem (Eikevliet)",51.0935,4.2835,24,"24 FTE (companyweb)","info@topfloor.be","","+32 3 889 38 15"),
("Vétoquinol Belgium","Farmaceutica veterinaria","Galileilaan 11/401","Niel",51.1170,4.3360,33,"33 FTE (companyweb)","info.be@vetoquinol.com","","+32 3 877 44 34"),
("Air Liquide Medical","Gas medicali","Tolhuisstraat 46-48","Schelle",51.1270,4.3340,150,"range 100–199 (Gouden Gids)","contact.be@airliquide.com","","+32 3 870 84 00"),
("Studio 100 (sede centrale)","Media/intrattenimento","Halfstraat 80","Schelle",51.1310,4.3480,169,"169 FTE (2025)","info@studio100.be","","+32 3 877 60 35"),
("Besix Infra","Costruzioni/infrastrutture","Steenwinkelstraat 640","Schelle",51.1190,4.3570,323,"323 FTE (companyweb)","besixinfra@besix.com","","+32 3 870 79 70"),
("Pfizer Manufacturing Belgium","Farmaceutica – produzione","Rijksweg 12","Puurs-Sint-Amands",51.0786,4.2640,4021,"4.021 FTE (companyweb) – maggior datore di lavoro della zona","n.d. (nessuna email generica pubblica)","","+32 3 890 92 11"),
("Novartis Manufacturing","Farmaceutica – produzione","Rijksweg 14","Puurs-Sint-Amands",51.0770,4.2600,980,"980 FTE (companyweb)","n.d. (nessuna email generica pubblica)","","+32 3 890 28 00"),
("Peleman Industries","Industria – rilegatura/stampa","Rijksweg 7","Puurs-Sint-Amands",51.0810,4.2680,84,"84 FTE (companyweb)","info.be@peleman.com","","+32 3 889 32 41"),
("Lonza Capsules (Capsugel Belgium)","Farmaceutica – capsule","Rijksweg 11","Bornem",51.0860,4.2560,474,"474 FTE (companyweb)","n.d. (nessuna email generica pubblica)","","+32 3 890 05 11"),
("Qualiphar","Farmaceutica/OTC","Rijksweg 9","Bornem",51.0870,4.2570,118,"118 dip. (fonte terza, da verificare)","receptie@qualiphar.com","","+32 3 889 17 21"),
("Arseus Medical","Forniture medicali","Rijksweg 10","Bornem",51.0870,4.2560,66,"66 FTE (Arseus Medical Group, companyweb)","info@arseus-medical.be","","+32 800 76 777"),
("European Master Batch (Sioen Chemicals)","Chimica – masterbatch","Rijksweg 15","Bornem",51.0850,4.2550,108,"108 FTE (companyweb)","info@sioenchemicals.com","","+32 3 890 64 00"),
("Smurfit Westrock Bornem","Packaging","Rijksweg 18","Bornem",51.0850,4.2540,151,"151 FTE (companyweb)","onthaal.bornem@multipkg.com","","+32 3 889 68 11"),
("LyondellBasell (A. Schulman Plastics)","Chimica – compound plastici","Pedro Colomalaan 25","Bornem",51.0900,4.2520,195,"195 FTE (companyweb)","n.d. (nessuna email generica pubblica)","","—"),
("AZ Rivierenland – campus Bornem","Ospedale","Kasteelstraat 23","Bornem",51.0975,4.2430,1500,"stima, tutti i campus (Bornem/Rumst/Willebroek)","info@azr.be","","+32 3 880 90 11"),
("Olympus Belgium","Dispositivi medici/ottica","Gansbroekstraat 70","Puurs-Sint-Amands",51.0750,4.2980,48,"48 FTE (companyweb)","n.d. (solo modulo web olympus.be)","","+32 3 870 58 00"),
("DSV Road + DSV Air & Sea","Logistica/spedizioni","Schoonmansveld 40","Puurs-Sint-Amands",51.0820,4.3110,480,"254 + 226 FTE (companyweb)","info.sea@be.dsv.com","","+32 3 611 06 00"),
("DSV Solutions Puurs","Logistica – magazzino","Schoonmansveld 34","Puurs-Sint-Amands",51.0825,4.3100,500,"range 500–999 (Gouden Gids, da verificare)","puurs.solutions@be.dsv.com","","+32 3 860 35 00"),
("Neovia Logistics","Logistica","Schoonmansveld 1","Puurs-Sint-Amands",51.0830,4.3090,189,"189 FTE (companyweb)","reception@neovialogistics.com","","+32 2 263 46 11"),
("Transport Robbyns","Trasporti","Schoonmansveld 26","Puurs-Sint-Amands",51.0822,4.3105,70,"70 FTE (companyweb)","planning@robbyns.be","","+32 3 866 04 21"),
("Groven+","Serramenti/costruzioni","Schoonmansveld 50","Puurs-Sint-Amands",51.0818,4.3115,50,"~50 dip.","home@grovenplus.be","","+32 3 877 00 44"),
("Hemiksem – Lamifil","Industria – cavi/fili","Frederic Sheidlaan","Hemiksem",51.1395,4.3330,224,"224 FTE (companyweb)","info@lamifil.be","","+32 3 870 06 11"),
("Wolf Oil Corporation","Lubrificanti","Georges Gilliotstraat 52","Hemiksem",51.1400,4.3380,75,"75 FTE (companyweb)","info@wolfoil.com","","+32 3 870 00 00"),
("Champion Chemicals","Lubrificanti","Georges Gilliotstraat 52","Hemiksem",51.1400,4.3385,85,"85 FTE (companyweb)","info@championlubes.com","","+32 3 870 00 20"),
("Pelican Rouge","Caffè/vending","Industrieweg 10a","Boom",51.0965,4.3560,344,"344 FTE (companyweb)","n.d. (nessuna email generica pubblica)","","—"),
("Bakker Belgium (Greenyard)","Ortofrutta – logistica","Industrieweg 16","Boom",51.0960,4.3570,151,"151 FTE (companyweb)","info@bakkerbelgium.be","","+32 3 808 84 00"),
("Atlas Copco Rental Europe","Noleggio attrezzature","Industrieweg 1F","Boom",51.0970,4.3550,123,"123 FTE (companyweb)","rental.belgium@atlascopco.com","","+32 3 401 67 67"),
("Cordeel (sede centrale)","Costruzioni","Frank Van Dyckelaan 15","Temse",51.1255,4.2200,199,"199 FTE Cordeel Zetel Temse (companyweb)","info@cordeel.eu","","+32 3 710 55 00"),
("C-Metal (gruppo Cordeel)","Carpenteria metallica","Frank Van Dyckelaan 19A","Temse",51.1255,4.2205,51,"51 FTE (companyweb)","info@c-metal.eu","","+32 3 710 55 62"),
("De Meyer NV","Meccanica/engineering","Frank Van Dyckelaan 28","Temse",51.1258,4.2210,130,"~130 dip. (sito aziendale)","info@demeyer.be","","+32 3 766 33 33"),
("BD – Becton Dickinson Distribution Center","Medtech – distribuzione","Laagstraat 57 (IZ TTS)","Temse",51.1340,4.2290,434,"434 FTE (companyweb)","info.benelux@bd.com","","+32 3 710 32 06"),
("Multi.Engineering","Engineering","Orlaylaan 10 (IZ TTS)","Temse",51.1345,4.2280,54,"54 FTE (companyweb)","info.temse@multi.engineering","","+32 3 710 58 10"),
("Modemakers Fashion","Moda – distribuzione","Kapelanielaan 17 (IZ TTS)","Temse",51.1350,4.2270,35,"35 FTE (companyweb)","info@modemakers.be","","+32 3 886 59 02"),
("Duvel Moortgat","Birrificio","Breendonk-Dorp 58","Puurs-Sint-Amands (Breendonk)",51.0450,4.3300,517,"~517 dip.","info@duvel.be","","+32 3 860 94 00"),
("Ceva Logistics Belgium","Logistica","Koningin Astridlaan 12","Willebroek",51.0680,4.3550,723,"723 FTE (companyweb)","receptie.willebroek@cevalogistics.com","","+32 3 860 45 00"),
("Distrilog Group","Logistica","Koningin Astridlaan 14","Willebroek",51.0675,4.3555,563,"563 FTE (companyweb)","info@distrilog.be","personeel@distrilog.be","+32 3 897 19 90"),
("Keppel Seghers Belgium","Engineering ambientale","Hoofd 1","Willebroek",51.0640,4.3620,117,"117 FTE (companyweb)","info.keppelseghers@keppel.com","","+32 3 880 77 00"),
("RCT Stevedoring","Logistica portuale","Boomsesteenweg 180","Willebroek",51.0750,4.3570,15,"15 FTE (companyweb)","info@rct-stevedoring.com","","+32 3 886 37 11"),
("Fero Group","Segnaletica/costruzioni","Jozef De Blockstraat 81","Willebroek",51.0580,4.3580,174,"174 FTE (companyweb)","info@ferogroup.be","","+32 3 288 75 28"),
("Colis Privé BeLux","Pacchi/logistica","Schoondonkweg 4","Willebroek",51.0510,4.3720,71,"71 FTE (companyweb)","infobelux@colisprive.com","","—"),
("PostNL – centro smistamento","Pacchi/logistica","MG Park De Hulst","Willebroek",51.0500,4.3750,400,"~400 a regime, incl. 250 autisti (stampa)","n.d. (nessuna email locale pubblica)","","—"),
("Etex Building Performance (Promat)","Materiali edili – R&D","Bormstraat 24","Willebroek (Tisselt)",51.0350,4.3550,273,"273 FTE (companyweb)","info@promat.be","","—"),
("Redwire Space","Aerospazio","Hogenakkerhoekstraat 9","Kruibeke",51.1660,4.3050,124,"124 FTE (companyweb)","n.d. (nessuna email generica pubblica)","","+32 3 250 14 14"),
("Verhaert","Innovazione/product design","Hogenakkerhoekstraat 21","Kruibeke",51.1662,4.3055,68,"68 FTE (companyweb)","info@verhaert.com","","+32 3 250 19 00"),
("Elscolab","Strumentazione laboratorio","Hogenakkerhoekstraat 14","Kruibeke",51.1661,4.3052,34,"34 FTE (companyweb)","contact@elscolab.com","","+32 3 250 15 70"),
("TTC Trade & Transport Corporation","Trasporti","Kasteleinsstraat 21","Kruibeke",51.1720,4.3100,233,"233 FTE (companyweb)","administratie@ttctransport.be","","+32 3 252 20 80"),
("Tesla Belgium","Automotive – sede/service","Boomsesteenweg 8","Aartselaar",51.1420,4.3800,231,"231 FTE (companyweb)","n.d. (nessuna email generica pubblica)","","—"),
("Toolstation Belgium (sede)","Retail utensili","Boomsesteenweg 58A","Aartselaar",51.1360,4.3770,84,"84 FTE (companyweb)","info@toolstation.be","","078 05 00 96"),
("Grundfos Bellux","Pompe – vendita","Boomsesteenweg 81-83","Aartselaar",51.1340,4.3760,75,"range 50–99","infobellux@grundfos.com","","+32 3 870 73 00"),
("Belcar","Noleggio/servizi auto","Bist 12","Aartselaar",51.1370,4.3880,86,"86 FTE (companyweb)","info@belcar.be","","+32 3 870 81 30"),
("Protime (gruppo SD Worx)","Software HR","Kontichsesteenweg 54","Aartselaar",51.1380,4.3990,376,"~376 dip. (gruppo)","info@protime.be","","+32 3 870 60 30"),
("ODTH – Opslag en Distributie Ter Haeghe","Logistica chimica","Doelhaagstraat 72","Rumst (Terhagen)",51.0870,4.3950,150,"150 FTE (companyweb)","info@odth.be","","+32 3 880 72 60"),
]
# fix name typo
D = [(("Lamifil",)+r[1:]) if r[0].startswith("Hemiksem") else r for r in D]

def hav(lat,lon):
    p=math.pi/180
    a=math.sin((lat-O_LAT)*p/2)**2+math.cos(O_LAT*p)*math.cos(lat*p)*math.sin((lon-O_LON)*p/2)**2
    return 2*6371*math.asin(math.sqrt(a))
D = [r for r in D if hav(r[4],r[5])<=10.0]
D.sort(key=lambda r: hav(r[4],r[5]))
DM=maps_dm([dest(r[2],r[3]) for r in D]) if KEY else None
if KEY: D=[r for _,r in sorted(zip(DM,D),key=lambda x:(x[0][0] is None,x[0][0] or 0))]; DM=sorted(DM,key=lambda x:(x[0] is None,x[0] or 0))

wb=Workbook(); ws=wb.active; ws.title="Aziende"
F="Arial"
hdr_fill=PatternFill("solid",fgColor="1F3864"); hf=Font(name=F,bold=True,color="FFFFFF",size=10)
thin=Side(style="thin",color="BFBFBF"); bd=Border(top=thin,bottom=thin,left=thin,right=thin)

ws["A1"]="Aziende >10 dipendenti entro 10 km da Arte Bianca – Egied De Jonghestraat 165, 2880 Bornem (Hingene)"
ws["A1"].font=Font(name=F,bold=True,size=13)
ws["A2"]="Partenza:"; ws["B2"]=ORIGIN
ws["A2"].font=Font(name=F,size=9,italic=True); ws["B2"].font=Font(name=F,size=9)
ws["E2"]="Aziende in elenco:"; ws["F2"]=len(D)
ws["G2"]="Con email generica:"; ws["H2"]=sum(1 for r in D if "@" in r[8])
for c in ("E2","F2","G2","H2"): ws[c].font=Font(name=F,size=9,bold=c in("F2","H2"))

H=["#","Azienda","Settore","Indirizzo","Comune","Distanza (km, linea d'aria)","Dipendenti (FTE/stima)","Nota dipendenti / fonte","Email generica","Email HR generica","Telefono"]+(["Km auto (Google Maps)","Minuti auto (Google Maps)"] if KEY else [])+["Percorso Google Maps"]
for i,h in enumerate(H,1):
    c=ws.cell(row=4,column=i,value=h); c.font=hf; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); c.border=bd
ws.row_dimensions[4].height=32
alt=PatternFill("solid",fgColor="F2F2F2")
for n,r in enumerate(D):
    row=5+n
    name,sett,ind,com,lat,lon,dip,nota,em,hr,tel=r
    dd=dest(ind,com)
    vals=[n+1,name,sett,ind,com,None,dip,nota,em,hr,tel]+(list(DM[n]) if KEY else [])+["Apri percorso"]
    for i,v in enumerate(vals,1):
        c=ws.cell(row=row,column=i,value=v); c.font=Font(name=F,size=10); c.border=bd
        c.alignment=Alignment(vertical="top",wrap_text=i in (2,3,8))
        if n%2: c.fill=alt
    ws.cell(row=row,column=6,value=round(hav(lat,lon),1))
    ws.cell(row=row,column=6).number_format="0.0"
    ws.cell(row=row,column=7).number_format="#,##0"
    if em.startswith("n.d."): ws.cell(row=row,column=9).font=Font(name=F,size=10,italic=True,color="C00000")
    lc=ws.cell(row=row,column=len(H)); lc.hyperlink=maps_link(dd); lc.font=Font(name=F,size=10,color="0563C1",underline="single")
widths=[4,34,24,26,22,12,12,34,36,22,16]+([12,12] if KEY else [])+[16]
for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes="C5"
ws.auto_filter.ref=f"A4:{get_column_letter(len(H))}{4+len(D)}"
last=4+len(D)
ws.cell(row=last+2,column=1,value="Note: distanza in linea d'aria da coordinate stimate (±0,5–1 km). Il link «Apri percorso» mostra su Google Maps km e tempo reali in auto (attenzione: Temse, Kruibeke e Hemiksem sono oltre la Schelda/Rupel, in auto il percorso è molto più lungo). Dipendenti = FTE dell'ultimo bilancio depositato (companyweb.be) salvo diversa indicazione. Email in rosso = nessun indirizzo generico pubblico: usare telefono/centralino o LinkedIn HR.").font=Font(name=F,size=9,italic=True)

# Template sheet
t=wb.create_sheet("Email template NL")
lines=[
("Oggetto / Onderwerp:","Business lunch op een paar minuten van {bedrijf} – Arte Bianca Hingene"),
("",""),
("Testo:","Beste HR-team,"),
("",""),
("","Wij zijn Arte Bianca, restaurant in de Egied De Jonghestraat 165 in Hingene (Bornem), op slechts {afstand} km van {bedrijf}."),
("",""),
("","Voor bedrijven uit de buurt hebben we een business lunch: [menu/gerecht] voor €[prijs], vlot geserveerd zodat uw medewerkers binnen het uur terug op het werk zijn. Ideaal voor teamlunches, klantafspraken of als attentie voor uw personeel."),
("",""),
("","Bekijk het menu: https://artebianca.be/business-lunch"),
("",""),
("","Voor vaste groepen bieden we graag bedrijfsvoorwaarden aan (reservatie op vaste dagen, facturatie op bedrijfsnaam)."),
("","Mogen we u vragen dit door te sturen naar de collega's die teamactiviteiten of personeelsvoordelen organiseren?"),
("",""),
("","Met vriendelijke groeten,"),
("","[Naam] – Arte Bianca"),
("","[telefoon] – artebianca.be"),
("",""),
("","Wenst u geen berichten meer van ons te ontvangen? Antwoord met 'uitschrijven'."),
("",""),
("Uso:","Sostituisci {bedrijf} e {afstand} con le colonne B e F del foglio Aziende (stampa unione Gmail/Word). Invia in blocchi da 20–30/giorno per evitare lo spam."),
("GDPR:","In Belgio le email B2B a indirizzi generici (info@, receptie@) sono ammesse senza consenso preventivo se pertinenti all'attività del destinatario e con possibilità di opt-out: lascia sempre la riga 'uitschrijven'."),
]
for i,(a,b) in enumerate(lines,1):
    t.cell(row=i,column=1,value=a).font=Font(name=F,bold=True,size=10)
    c=t.cell(row=i,column=2,value=b); c.font=Font(name=F,size=10); c.alignment=Alignment(wrap_text=True,vertical="top")
t.column_dimensions["A"].width=20; t.column_dimensions["B"].width=100

out="/home/user/Spauz23/marketing/aziende_10km_arte_bianca.xlsx"
wb.save(out); print(out,len(D))
for r in D: print(round(hav(r[4],r[5]),1), r[0])
